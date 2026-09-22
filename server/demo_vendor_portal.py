import http.server
import socketserver
import urllib.parse
import io
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

PORT = 8085

class DemoVendorPortalHandler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(200)
        self.end_headers()

    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        content_len = int(self.headers.get('Content-Length', 0))
        body = self.rfile.read(content_len).decode('utf-8', errors='ignore')
        form_data = urllib.parse.parse_qs(body)

        if path == '/demo-vendor/api/login':
            username = form_data.get('username', [''])[0]
            password = form_data.get('password', [''])[0]

            # Scenario testing: support failure demonstration
            if password == 'invalid_password' or password == 'wrong':
                self.send_response(401)
                self.send_header('Content-Type', 'application/json')
                self.end_headers()
                self.wfile.write(json.dumps({"success": False, "error": "LOGIN_FAILED", "message": "Invalid username or password"}).encode('utf-8'))
                return

            self.send_response(200)
            self.send_header('Content-Type', 'application/json')
            self.end_headers()
            self.wfile.write(json.dumps({
                "success": True,
                "sessionToken": "TB_DEMO_SESS_89214710",
                "redirectUrl": "/demo-vendor/dashboard"
            }).encode('utf-8'))
            return

        self.send_response(404)
        self.end_headers()

    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        query_params = urllib.parse.parse_qs(parsed.query)

        # Failure scenario switches
        scenario = query_params.get('scenario', ['normal'])[0]

        if path == '/' or path == '/demo-vendor' or path == '/demo-vendor/login':
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()

            user_field_id = "username" if scenario != "login_element_changed" else "user_email_modified"
            user_testid = "username-input" if scenario != "login_element_changed" else "user-email-legacy"

            html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>TransBingo Partner Portal - School Fee Collection</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #071A36; color: #fff; display: flex; align-items: center; justify-content: center; height: 100vh; margin: 0; }}
        .card {{ background: #0F2A4A; padding: 36px; border-radius: 12px; width: 380px; box-shadow: 0 10px 30px rgba(0,0,0,0.5); border: 1px solid #1E3E62; }}
        h1 {{ font-size: 20px; color: #fff; margin-top: 0; margin-bottom: 4px; font-weight: 700; }}
        .sub {{ font-size: 12px; color: #8FA2B7; margin-bottom: 24px; }}
        label {{ display: block; font-size: 11px; font-weight: 600; text-transform: uppercase; color: #8FA2B7; margin-bottom: 6px; }}
        input {{ width: 100%; padding: 10px 12px; background: #071A36; border: 1px solid #1E3E62; border-radius: 8px; color: #fff; font-size: 13px; margin-bottom: 16px; box-sizing: border-box; }}
        input:focus {{ outline: none; border-color: #1267E8; }}
        button {{ width: 100%; padding: 11px; background: #1267E8; color: white; border: none; border-radius: 8px; font-size: 13px; font-weight: 700; cursor: pointer; }}
        button:hover {{ background: #0E52BA; }}
        .meta {{ font-size: 11px; color: #5B728D; text-align: center; margin-top: 20px; }}
    </style>
</head>
<body>
    <div class="card" id="login-card" data-testid="login-container">
        <h1>TransBingo Portal</h1>
        <div class="sub">Institutional Fee Collection & Settlement</div>
        <form id="login-form" action="/demo-vendor/dashboard" method="GET">
            <label>Username / Partner ID</label>
            <input type="text" id="{user_field_id}" name="username" data-testid="{user_testid}" value="demo-operator" required />
            
            <label>Password</label>
            <input type="password" id="password" name="password" data-testid="password-input" value="••••••••••••" required />
            
            <button type="submit" id="btn-login" data-testid="login-btn">Sign In to Portal</button>
        </form>
        <div class="meta">TransBingo Demo Environment • API v4.2</div>
    </div>
</body>
</html>"""
            self.wfile.write(html.encode('utf-8'))
            return

        elif path == '/demo-vendor/dashboard':
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()

            report_menu_text = "Collection Report"
            report_menu_id = "menu-collection-report"
            if scenario == "report_menu_changed":
                report_menu_text = "Payment Collection Report"
                report_menu_id = "menu-payment-collection-report"

            html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>TransBingo - Operator Dashboard</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #F5F7FA; color: #14213D; margin: 0; }}
        .header {{ background: #071A36; color: white; padding: 16px 30px; display: flex; justify-content: space-between; align-items: center; }}
        .nav {{ background: #0F2A4A; color: white; padding: 10px 30px; display: flex; gap: 20px; font-size: 13px; font-weight: 600; }}
        .nav a {{ color: #CBD5E1; text-decoration: none; padding: 6px 12px; border-radius: 6px; }}
        .nav a:hover, .nav a.active {{ background: #1267E8; color: white; }}
        .content {{ padding: 30px; max-width: 900px; margin: 0 auto; }}
        .box {{ background: white; border: 1px solid #E2E8F0; border-radius: 12px; padding: 24px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }}
    </style>
</head>
<body>
    <div class="header" data-testid="dashboard-view" id="dashboard-view">
        <div style="font-weight:bold;font-size:16px;">TransBingo Partner Console</div>
        <div style="font-size:12px;color:#94A3B8;">Operator: demo-operator</div>
    </div>
    <div class="nav" id="main-nav">
        <a href="/demo-vendor/dashboard" class="active">Overview</a>
        <a href="/demo-vendor/reports" id="menu-reports" data-testid="menu-reports">Reports</a>
        <a href="/demo-vendor/reports?scenario={scenario}" id="{report_menu_id}" data-testid="{report_menu_id}">{report_menu_text}</a>
    </div>
    <div class="content">
        <div class="box">
            <h2>Welcome, demo-operator</h2>
            <p style="font-size:13px;color:#64748B;">Navigate to <strong>Reports &gt; {report_menu_text}</strong> to extract institutional settlements.</p>
            <a href="/demo-vendor/reports?scenario={scenario}" style="display:inline-block;padding:10px 18px;background:#1267E8;color:white;text-decoration:none;border-radius:8px;font-size:13px;font-weight:600;">
                Open {report_menu_text} &rarr;
            </a>
        </div>
    </div>
</body>
</html>"""
            self.wfile.write(html.encode('utf-8'))
            return

        elif path == '/demo-vendor/reports':
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()

            school_selected = query_params.get('schoolCode', ['UTTARA_MDL'])[0]
            date_selected = query_params.get('date', ['2026-09-18'])[0]
            is_searched = 'searched' in query_params or 'schoolCode' in query_params

            html = f"""<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>TransBingo - Collection Report</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #F5F7FA; color: #14213D; margin: 0; }}
        .header {{ background: #071A36; color: white; padding: 14px 30px; display: flex; justify-content: space-between; align-items: center; }}
        .content {{ padding: 30px; max-width: 1000px; margin: 0 auto; }}
        .box {{ background: white; border: 1px solid #E2E8F0; border-radius: 12px; padding: 24px; box-shadow: 0 1px 3px rgba(0,0,0,0.05); }}
        .form-grid {{ display: grid; grid-template-columns: 2fr 1fr 1fr 1fr; gap: 15px; align-items: flex-end; margin-bottom: 24px; }}
        label {{ display: block; font-size: 11px; font-weight: 700; text-transform: uppercase; color: #64748B; margin-bottom: 6px; }}
        select, input {{ width: 100%; padding: 9px 12px; border: 1px solid #CBD5E1; border-radius: 8px; font-size: 13px; box-sizing: border-box; }}
        button.btn-search {{ padding: 10px 18px; background: #1267E8; color: white; border: none; border-radius: 8px; font-size: 13px; font-weight: 700; cursor: pointer; }}
        button.btn-search:hover {{ background: #0E52BA; }}
        table {{ width: 100%; border-collapse: collapse; font-size: 12px; margin-top: 15px; }}
        th, td {{ padding: 10px 14px; text-align: left; border-bottom: 1px solid #E2E8F0; }}
        th {{ background: #F8FAFC; color: #475569; font-weight: 600; text-transform: uppercase; font-size: 11px; }}
        .btn-export {{ display: inline-flex; align-items: center; gap: 8px; padding: 10px 18px; background: #16A34A; color: white; text-decoration: none; border-radius: 8px; font-size: 13px; font-weight: 700; }}
        .btn-export:hover {{ background: #15803D; }}
    </style>
</head>
<body>
    <div class="header">
        <div style="font-weight:bold;font-size:15px;">TransBingo &bull; Institutional Collection Report</div>
        <div style="font-size:12px;color:#94A3B8;">Environment: DEMO PORTAL</div>
    </div>
    <div class="content">
        <div class="box">
            <h2 style="margin-top:0;font-size:18px;margin-bottom:18px;">Search Settlement Transactions</h2>
            <form id="filter-form" action="/demo-vendor/reports" method="GET">
                <input type="hidden" name="searched" value="true" />
                <input type="hidden" name="scenario" value="{scenario}" />
                <div class="form-grid">
                    <div>
                        <label>Target School</label>
                        <select id="school-select" name="schoolCode" data-testid="school-select">
                            <option value="UTTARA_MDL" {"selected" if school_selected == "UTTARA_MDL" else ""}>Uttara Model High School</option>
                            <option value="ABC_INT" {"selected" if school_selected == "ABC_INT" else ""}>ABC School</option>
                            <option value="DHAKA_MDL" {"selected" if school_selected == "DHAKA_MDL" else ""}>Dhaka Model School</option>
                        </select>
                    </div>
                    <div>
                        <label>From Date</label>
                        <input type="date" id="from-date" name="fromDate" data-testid="from-date" value="{date_selected}" />
                    </div>
                    <div>
                        <label>To Date</label>
                        <input type="date" id="to-date" name="toDate" data-testid="to-date" value="{date_selected}" />
                    </div>
                    <div>
                        <button type="submit" id="search-btn" class="btn-search" data-testid="search-btn">Search Records</button>
                    </div>
                </div>
            </form>

            <div id="report-table" data-testid="report-table" style="margin-top:20px;">
                <div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:12px;">
                    <div>
                        <span style="font-weight:bold;font-size:14px;">Results for {school_selected} ({date_selected})</span>
                        <span style="font-size:12px;color:#64748B;margin-left:8px;">1,842 transactions found &bull; Total: ৳4,821,500</span>
                    </div>
                    <a href="/demo-vendor/export.xlsx?schoolCode={school_selected}&date={date_selected}&scenario={scenario}" id="export-excel-btn" class="btn-export" data-testid="export-excel-btn">
                        <span>&#128196;</span> Export Excel (.xlsx)
                    </a>
                </div>

                <table>
                    <thead>
                        <tr>
                            <th>Transaction ID</th>
                            <th>Student ID</th>
                            <th>Student Name</th>
                            <th>Amount (BDT)</th>
                            <th>Payment Date</th>
                            <th>Method</th>
                            <th>Status</th>
                        </tr>
                    </thead>
                    <tbody>
                        <tr>
                            <td style="font-family:monospace;font-weight:bold;color:#1267E8;">TXN001</td>
                            <td style="font-family:monospace;">STU1001</td>
                            <td>Tanvir Rahman</td>
                            <td style="font-weight:bold;">৳500</td>
                            <td>18/09/2026</td>
                            <td>bKash Online</td>
                            <td><span style="color:#16A34A;font-weight:bold;">PAID</span></td>
                        </tr>
                        <tr>
                            <td style="font-family:monospace;font-weight:bold;color:#1267E8;">TXN002</td>
                            <td style="font-family:monospace;">STU1002</td>
                            <td>Sadia Sultana</td>
                            <td style="font-weight:bold;">৳1,000</td>
                            <td>18/09/2026</td>
                            <td>Nagad Checkout</td>
                            <td><span style="color:#16A34A;font-weight:bold;">PAID</span></td>
                        </tr>
                        <tr>
                            <td style="font-family:monospace;font-weight:bold;color:#1267E8;">TXN003</td>
                            <td style="font-family:monospace;">STU1003</td>
                            <td>Nabil Ahmed</td>
                            <td style="font-weight:bold;">৳5,500</td>
                            <td>18/09/2026</td>
                            <td>DBBL Rocket</td>
                            <td><span style="color:#16A34A;font-weight:bold;">PAID</span></td>
                        </tr>
                        <tr>
                            <td style="font-family:monospace;font-weight:bold;color:#1267E8;">TXN004</td>
                            <td style="font-family:monospace;">STU1004</td>
                            <td>Farhana Kabir</td>
                            <td style="font-weight:bold;">৳4,500</td>
                            <td>18/09/2026</td>
                            <td>Visa / Card</td>
                            <td><span style="color:#16A34A;font-weight:bold;">PAID</span></td>
                        </tr>
                    </tbody>
                </table>
                <div style="font-size:11px;color:#94A3B8;margin-top:12px;text-align:right;">Showing top 4 of 1,842 records (Click Export Excel for complete batch).</div>
            </div>
        </div>
    </div>
</body>
</html>"""
            self.wfile.write(html.encode('utf-8'))
            return

        elif path == '/demo-vendor/export.xlsx':
            if scenario == "download_timeout":
                import time
                time.sleep(12)  # Trigger timeout
                return

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Collection_Report"

            # Header styling
            header_fill = PatternFill(start_color="071A36", end_color="071A36", fill_type="solid")
            header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            headers = ["Row", "Reference", "StudentID", "StudentName", "Amount", "Timestamp", "Status", "PaymentMethod", "SchoolCode"]
            ws.append(headers)

            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Deterministic rows matching Section 12 specifications:
            # File: TransBingo_Collection_20260918.xlsx
            # Rows: 1,842
            # Total Amount: ৳4,821,500
            total_rows = 1842
            target_sum = 4821500

            # Seed top 4 rows
            sample_rows = [
                (1, "TXN001", "STU1001", "Tanvir Rahman", 500, "18/09/2026 08:30:12", "PAID", "bKash Online", "UTTARA_MDL"),
                (2, "TXN002", "STU1002", "Sadia Sultana", 1000, "18/09/2026 08:34:20", "PAID", "Nagad Checkout", "UTTARA_MDL"),
                (3, "TXN003", "STU1003", "Nabil Ahmed", 5500, "18/09/2026 09:12:05", "PAID", "DBBL Rocket", "UTTARA_MDL"),
                (4, "TXN004", "STU1004", "Farhana Kabir", 4500, "18/09/2026 09:15:33", "PAID", "Visa / Card", "UTTARA_MDL"),
            ]
            for r in sample_rows:
                ws.append(list(r))

            current_sum = sum(r[4] for r in sample_rows)
            remaining_rows = total_rows - len(sample_rows)
            base_amt = 2600

            for i in range(len(sample_rows) + 1, total_rows + 1):
                if i == total_rows:
                    amt = target_sum - current_sum
                else:
                    amt = 2600 if i % 2 == 0 else 2620
                    current_sum += amt
                ref = f"TXN{i:04d}"
                std_id = f"STU{1000 + i}"
                ws.append([i, ref, std_id, f"Student {i}", amt, "18/09/2026 10:00:00", "PAID", "bKash", "UTTARA_MDL"])

            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            xlsx_bytes = output.getvalue()

            self.send_response(200)
            self.send_header('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            self.send_header('Content-Disposition', 'attachment; filename="TransBingo_Collection_20260918.xlsx"')
            self.send_header('Content-Length', str(len(xlsx_bytes)))
            self.end_headers()
            self.wfile.write(xlsx_bytes)
            return

        self.send_response(404)
        self.end_headers()
        self.wfile.write(b"Not Found")

def run():
    socketserver.TCPServer.allow_reuse_address = True
    with socketserver.TCPServer(("", PORT), DemoVendorPortalHandler) as httpd:
        print(f"Demo Vendor Portal active on http://localhost:{PORT}")
        httpd.serve_forever()

if __name__ == "__main__":
    run()
